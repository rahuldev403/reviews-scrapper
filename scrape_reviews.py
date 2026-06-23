"""
Scrape ALL 1-star reviews of a Google Play app and export them to an Excel sheet.

How it works
------------
The Play Store website doesn't ship reviews in its HTML; it loads them on demand
from an internal API. The `google-play-scraper` library talks to that same API,
so we get clean, structured data instead of parsing messy HTML.

We ask the API for 1-star reviews only (`filter_score_with=1`) and walk through
the results page by page using a "continuation token" the API hands back, until
there are no more pages left.

Usage
-----
    python scrape_reviews.py
    python scrape_reviews.py --app-id org.grow90.whatsub --score 1 --country in --lang en
"""

import argparse
import sys
import time
from datetime import datetime

from google_play_scraper import Sort, reviews


def parse_args():
    parser = argparse.ArgumentParser(
        description="Scrape Google Play reviews of a given star rating into Excel."
    )
    parser.add_argument(
        "--app-id",
        default="org.grow90.whatsub",
        help="Play Store application id (the ?id= value in the URL).",
    )
    parser.add_argument(
        "--score",
        type=int,
        default=1,
        choices=[1, 2, 3, 4, 5],
        help="Which star rating to scrape (default: 1).",
    )
    parser.add_argument("--country", default="in", help="Country code (default: in).")
    parser.add_argument("--lang", default="en", help="Language code (default: en).")
    parser.add_argument(
        "--page-size",
        type=int,
        default=200,
        help="Reviews requested per API call (max ~200).",
    )
    parser.add_argument(
        "--output",
        default=None,
        help="Output .xlsx path. Defaults to reviews_<appid>_<score>star_<timestamp>.xlsx",
    )
    return parser.parse_args()


def fetch_all_reviews(app_id, score, country, lang, page_size):
    """Page through the reviews API collecting every review of `score` stars."""
    all_reviews = []
    seen_ids = set()
    continuation_token = None
    page = 0

    print(f"\nScraping {score}-star reviews for '{app_id}' "
          f"(country={country}, lang={lang})\n")

    while True:
        page += 1
        try:
            batch, continuation_token = reviews(
                app_id,
                lang=lang,
                country=country,
                sort=Sort.NEWEST,
                count=page_size,
                filter_score_with=score,
                continuation_token=continuation_token,
            )
        except Exception as exc:  # noqa: BLE001 - surface any API/network error to the user
            print(f"  ! Error on page {page}: {exc}")
            print("    Stopping early; keeping what we have so far.")
            break

        # Deduplicate (the API can occasionally repeat a review across pages).
        new_in_batch = 0
        for r in batch:
            rid = r.get("reviewId")
            if rid in seen_ids:
                continue
            seen_ids.add(rid)
            all_reviews.append(r)
            new_in_batch += 1

        print(f"  page {page:>3}: got {len(batch):>3} reviews "
              f"({new_in_batch} new)  ->  total so far: {len(all_reviews)}")

        # The API signals "no more pages" by returning an empty token / no new data.
        if not continuation_token or not batch or new_in_batch == 0:
            break

        time.sleep(0.4)  # be polite to the endpoint

    print(f"\nDone. Collected {len(all_reviews)} unique {score}-star reviews.\n")
    return all_reviews


def to_dataframe(raw_reviews):
    import pandas as pd

    rows = []
    for r in raw_reviews:
        rows.append(
            {
                "Review ID": r.get("reviewId"),
                "User Name": r.get("userName"),
                "Rating": r.get("score"),
                "Review": r.get("content"),
                "Thumbs Up": r.get("thumbsUpCount"),
                "App Version": r.get("reviewCreatedVersion"),
                "Date": r.get("at"),
                "Developer Reply": r.get("replyContent"),
                "Reply Date": r.get("repliedAt"),
            }
        )
    df = pd.DataFrame(rows)
    # Newest first, and make the date timezone-naive so Excel accepts it.
    if not df.empty:
        df = df.sort_values("Date", ascending=False, na_position="last")
        for col in ("Date", "Reply Date"):
            df[col] = pd.to_datetime(df[col], errors="coerce").dt.tz_localize(None)
    return df


def write_excel(df, output_path, score):
    from openpyxl.utils import get_column_letter

    sheet_name = f"{score}-star reviews"
    with __import__("pandas").ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        ws = writer.sheets[sheet_name]

        # Freeze the header row and turn on autofilter.
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        # Reasonable column widths; wrap the long "Review" / "Reply" columns.
        widths = {
            "Review ID": 22,
            "User Name": 22,
            "Rating": 8,
            "Review": 70,
            "Thumbs Up": 11,
            "App Version": 13,
            "Date": 20,
            "Developer Reply": 50,
            "Reply Date": 20,
        }
        from openpyxl.styles import Alignment, Font

        header_font = Font(bold=True)
        for idx, col_name in enumerate(df.columns, start=1):
            letter = get_column_letter(idx)
            ws.column_dimensions[letter].width = widths.get(col_name, 18)
            ws.cell(row=1, column=idx).font = header_font
            if col_name in ("Review", "Developer Reply"):
                for row in range(2, ws.max_row + 1):
                    ws.cell(row=row, column=idx).alignment = Alignment(wrap_text=True, vertical="top")

    print(f"Excel written: {output_path}")


def main():
    args = parse_args()

    raw = fetch_all_reviews(
        app_id=args.app_id,
        score=args.score,
        country=args.country,
        lang=args.lang,
        page_size=args.page_size,
    )

    df = to_dataframe(raw)

    output = args.output
    if output is None:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_app = args.app_id.replace(".", "_")
        output = f"reviews_{safe_app}_{args.score}star_{stamp}.xlsx"

    if df.empty:
        print("No reviews found; nothing to write.")
        sys.exit(0)

    write_excel(df, output, args.score)
    print(f"\nSummary: {len(df)} rows exported.")


if __name__ == "__main__":
    main()
