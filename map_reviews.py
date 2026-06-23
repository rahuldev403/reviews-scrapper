"""
Map the 258 scraped 1-star reviews against the team's tracking sheets and produce
a NEW sheet that is the scraped sheet PLUS columns showing, for each review:

    Tracked     -> Yes / No  (is this review present in a tracking sheet?)
    Status      -> done / pending / agreed / no reply / ...  (from the tracker)
    Tracked In  -> New / Sheet1  (which tracking sheet it was found in)
    Match Type  -> exact / fuzzy / name  (HOW it was linked, for transparency)

The base rows and original columns of the scraped sheet are kept unchanged.
The heavier tracking fields (phone, name in App, assigned, Remarks, ...) are
deliberately NOT carried over.

Run:
    python map_reviews.py
"""

import re
from datetime import datetime
from difflib import SequenceMatcher, get_close_matches

import pandas as pd

WORKBOOK = "Playstore Ratings.xlsx"
SCRAPED_SHEET = "1-star reviews"
TRACKING_SHEETS = ["New", "Sheet1"]  # both have `review` + `status`
FUZZY_CUTOFF = 0.87  # 0..1 ; higher = stricter text match


def norm(s):
    """Lowercase, strip punctuation/whitespace for robust comparison."""
    if pd.isna(s):
        return ""
    s = str(s).lower().strip()
    s = re.sub(r"[^a-z0-9 ]", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def build_tracking_index():
    """
    Returns:
      text_index: {normalized_review_text: (status, source)}
      name_index: {normalized_name: (status, source)}  -- only names that are
                  unambiguous (appear once) are kept, to avoid false links.
    """
    text_index = {}
    name_counts = {}
    name_tmp = {}

    for sheet in TRACKING_SHEETS:
        df = pd.read_excel(WORKBOOK, sheet_name=sheet)
        name_col = "name in Playstore" if "name in Playstore" in df.columns else "name"
        for _, row in df.iterrows():
            status = row.get("status")
            status = "" if pd.isna(status) else str(status).strip()

            key = norm(row.get("review"))
            if len(key) > 8:
                # Prefer the first non-blank status; prefer New (processed first).
                if key not in text_index or (not text_index[key][0] and status):
                    text_index[key] = (status, sheet)

            nkey = norm(row.get(name_col))
            if len(nkey) >= 4:
                name_counts[nkey] = name_counts.get(nkey, 0) + 1
                if nkey not in name_tmp or (not name_tmp[nkey][0] and status):
                    name_tmp[nkey] = (status, sheet)

    # keep only unambiguous names
    name_index = {k: v for k, v in name_tmp.items() if name_counts[k] == 1}
    return text_index, name_index


def map_reviews():
    scraped = pd.read_excel(WORKBOOK, sheet_name=SCRAPED_SHEET)
    text_index, name_index = build_tracking_index()
    tracking_keys = list(text_index.keys())

    tracked, status_col, tracked_in, match_type = [], [], [], []

    for _, row in scraped.iterrows():
        rkey = norm(row.get("Review"))
        nkey = norm(row.get("User Name"))

        status, source, mtype = "", "", ""

        # 1) exact review-text match
        if len(rkey) > 8 and rkey in text_index:
            status, source = text_index[rkey]
            mtype = "exact"
        # 2) fuzzy review-text match
        elif len(rkey) > 8:
            close = get_close_matches(rkey, tracking_keys, n=1, cutoff=FUZZY_CUTOFF)
            if close:
                # double-check the ratio is genuinely high
                ratio = SequenceMatcher(None, rkey, close[0]).ratio()
                if ratio >= FUZZY_CUTOFF:
                    status, source = text_index[close[0]]
                    mtype = "fuzzy"
        # 3) fallback: unambiguous name match (only if text didn't link)
        if not mtype and len(nkey) >= 4 and nkey in name_index:
            status, source = name_index[nkey]
            mtype = "name"

        tracked.append("Yes" if mtype else "No")
        status_col.append(status.title() if status else "")
        tracked_in.append(source)
        match_type.append(mtype)

    scraped["Tracked"] = tracked
    scraped["Status"] = status_col
    scraped["Tracked In"] = tracked_in
    scraped["Match Type"] = match_type
    return scraped


def write_excel(df, path):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheet_name = "1-star reviews (mapped)"
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        ws = writer.sheets[sheet_name]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        widths = {
            "Review ID": 22, "User Name": 22, "Rating": 7, "Review": 65,
            "Thumbs Up": 10, "App Version": 12, "Date": 19,
            "Developer Reply": 45, "Reply Date": 19,
            "Tracked": 9, "Status": 14, "Tracked In": 11, "Match Type": 11,
        }
        header_font = Font(bold=True)
        wrap_cols = {"Review", "Developer Reply"}
        green = PatternFill("solid", fgColor="C6EFCE")  # tracked
        grey = PatternFill("solid", fgColor="F2F2F2")   # not tracked
        yellow = PatternFill("solid", fgColor="FFF2CC") # fuzzy/name -> review me

        col_idx = {name: i + 1 for i, name in enumerate(df.columns)}
        for name, i in col_idx.items():
            ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 16)
            ws.cell(row=1, column=i).font = header_font
            if name in wrap_cols:
                for r in range(2, ws.max_row + 1):
                    ws.cell(row=r, column=i).alignment = Alignment(wrap_text=True, vertical="top")

        # colour the Tracked column; flag fuzzy/name matches
        t_i, m_i = col_idx["Tracked"], col_idx["Match Type"]
        for r in range(2, ws.max_row + 1):
            is_tracked = ws.cell(row=r, column=t_i).value == "Yes"
            ws.cell(row=r, column=t_i).fill = green if is_tracked else grey
            if ws.cell(row=r, column=m_i).value in ("fuzzy", "name"):
                ws.cell(row=r, column=m_i).fill = yellow


def main():
    df = map_reviews()
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out = f"reviews_1star_tracked_{stamp}.xlsx"
    write_excel(df, out)

    total = len(df)
    tracked = (df["Tracked"] == "Yes").sum()
    print(f"\nTotal scraped 1-star reviews : {total}")
    print(f"Tracked (found in trackers)  : {tracked}")
    print(f"Not tracked (new/untracked)  : {total - tracked}")
    print("\nMatch type breakdown:")
    print(df[df["Match Type"] != ""]["Match Type"].value_counts().to_string())
    print("\nStatus breakdown (tracked rows):")
    sb = df[df["Status"] != ""]["Status"].value_counts()
    print(sb.to_string() if not sb.empty else "(none)")
    print(f"\nWritten: {out}")


if __name__ == "__main__":
    main()
